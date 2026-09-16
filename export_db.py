#!/usr/bin/env python3
"""Back up / export the zone-visit database to CSV and XLSX.

Standalone on purpose: it imports nothing from the tracker, so you can run it
while the app is running, or long after it has stopped. The database is opened
READ-ONLY through a file: URI, so this can never lock or corrupt the file the
tracker is writing to.

The usual reason to run it is a RunPod pod about to be stopped - only
/workspace survives, and an export you can download is worth more than a
database file on a dead pod.

  python export_db.py
  python export_db.py --db /workspace/footfall.db --out-dir /workspace/exports

Output files are timestamped, so running it repeatedly never overwrites an
earlier export:

  /workspace/exports/footfall_2026-09-11_1530.csv
  /workspace/exports/footfall_2026-09-11_1530.xlsx

A NOTE ON person_id, which colours everything below:

    person_id comes from the tracker and is stable only while a person's track
    is unbroken. It resets when someone is occluded, leaves and comes back, or
    is simply missed for long enough. One human walking through the shop can
    therefore produce several person_ids.

    So "unique visitors" here is an ESTIMATE, and it errs high - never a true
    headcount. Treat it as a trend line across comparable days, not a number to
    report as fact.
"""

from __future__ import annotations

import argparse
import csv
import os
import sqlite3
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

TABLE = "zone_visits"

# /workspace is the RunPod pod path this was written for, and it is still the
# right default there. On Windows it resolves to the root of the current drive
# (\workspace\footfall.db), which is neither writable nor where anyone put a
# database - so fall back to beside this script when /workspace does not exist.
_ON_POD = os.path.isdir("/workspace")
_HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DB = "/workspace/footfall.db" if _ON_POD else os.path.join(_HERE, "footfall.db")
DEFAULT_OUT_DIR = ("/workspace/exports" if _ON_POD
                   else os.path.join(_HERE, "exports"))

# Excel refuses any cell over 32767 characters; a long journey path can get
# there on a busy day, so paths are truncated well short of it.
MAX_PATH_CHARS = 2000

ESTIMATE_NOTE = (
    "person_id resets when a track breaks (occlusion, re-entry), so unique "
    "visitor counts are an estimate that errs high, not an exact headcount."
)


class ExportError(Exception):
    """Something the user can act on - reported as a message, not a traceback."""


# ---------------------------------------------------------------------------
# reading


def open_readonly(db_path: Path) -> sqlite3.Connection:
    """Open the database read-only.

    mode=ro means SQLite will not create the file, will not take a write lock,
    and will refuse any statement that would modify it - the tracker can keep
    writing while this runs.
    """
    if not db_path.exists():
        raise ExportError(
            f"No database at {db_path}\n"
            "  The tracker only writes one when you ask it to:\n"
            f"    python app.py --db {db_path}\n"
            "  Point --db here at that same file. On a RunPod pod it normally\n"
            "  lives at /workspace/footfall.db, since only /workspace survives."
        )
    if db_path.is_dir():
        raise ExportError(f"{db_path} is a directory, not a database file.")

    # as_uri() handles Windows drive letters and percent-encodes spaces, which
    # a hand-built "file:" + str(path) does not.
    uri = f"{db_path.resolve().as_uri()}?mode=ro"
    try:
        conn = sqlite3.connect(uri, uri=True)
    except sqlite3.OperationalError as exc:
        raise ExportError(f"Could not open {db_path} read-only: {exc}") from exc
    conn.row_factory = sqlite3.Row
    return conn


def read_rows(conn: sqlite3.Connection) -> tuple[List[str], List[sqlite3.Row]]:
    """Every row of the visits table, oldest first."""
    exists = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (TABLE,)
    ).fetchone()
    if not exists:
        found = [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        listing = ", ".join(found) if found else "none"
        raise ExportError(
            f"The database has no '{TABLE}' table (tables found: {listing}).\n"
            "  Either this is not the tracker's database, or the tracker has\n"
            "  not written a completed visit yet."
        )

    columns = [r["name"] for r in conn.execute(f"PRAGMA table_info({TABLE})")]
    # Order by entered_at so the raw dump and the journey paths agree; id is a
    # tiebreaker for rows written within the same second.
    order = "entered_at, id" if {"entered_at", "id"} <= set(columns) else "rowid"
    rows = conn.execute(f"SELECT * FROM {TABLE} ORDER BY {order}").fetchall()
    if not rows:
        raise ExportError(
            f"The '{TABLE}' table is empty - there is nothing to export yet.\n"
            "  A row is written when someone leaves a zone, so this is normal\n"
            "  if the tracker has only just started."
        )
    return columns, rows


def parse_ts(value) -> Optional[datetime]:
    """ISO-ish timestamp to datetime, or None if it can't be read.

    The tracker writes ISO strings, but an export should never die on one odd
    row, so anything unparseable is simply left out of the time-based tables.
    """
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip().replace("T", " ")
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# derived tables


def build_journeys(rows: Sequence[sqlite3.Row]) -> List[list]:
    """One row per person_id: when they appeared, how long, and where they went."""
    by_person: "OrderedDict[object, List[sqlite3.Row]]" = OrderedDict()
    for r in rows:
        by_person.setdefault(r["person_id"], []).append(r)

    out = []
    for pid, visits in by_person.items():
        entered = [t for t in (parse_ts(v["entered_at"]) for v in visits) if t]
        left = [t for t in (parse_ts(v["left_at"]) for v in visits) if t]
        total = sum(float(v["seconds"] or 0) for v in visits)
        zones = [str(v["zone_id"]) for v in visits]

        # Collapse immediate repeats. A single stay can be split into several
        # rows when the tracker loses someone for longer than its grace period
        # and re-opens the visit, which would otherwise render as
        # "aisle_1 > aisle_1 > aisle_1" and tell the reader nothing.
        path_parts: List[str] = []
        for z in zones:
            if not path_parts or path_parts[-1] != z:
                path_parts.append(z)
        path = " > ".join(path_parts)
        if len(path) > MAX_PATH_CHARS:
            keep = path[:MAX_PATH_CHARS].rsplit(" > ", 1)[0]
            dropped = len(path_parts) - keep.count(" > ") - 1
            path = f"{keep} > ... (+{dropped} more)"

        out.append([
            pid,
            min(entered).isoformat(sep=" ") if entered else "",
            max(left).isoformat(sep=" ") if left else "",
            round(total, 1),
            len(set(zones)),
            len(visits),
            path,
        ])
    return out


def build_summary(rows: Sequence[sqlite3.Row]) -> dict:
    """Per-hour unique visitors, per-zone totals, and the busiest hour."""
    per_hour_people: "defaultdict[str, set]" = defaultdict(set)
    per_hour_visits: Counter = Counter()
    per_zone_visits: Counter = Counter()
    per_zone_seconds: "defaultdict[str, float]" = defaultdict(float)
    undated = 0

    for r in rows:
        zone = str(r["zone_id"])
        per_zone_visits[zone] += 1
        per_zone_seconds[zone] += float(r["seconds"] or 0)

        ts = parse_ts(r["entered_at"])
        if ts is None:
            undated += 1
            continue
        bucket = ts.strftime("%Y-%m-%d %H:00")
        per_hour_people[bucket].add(r["person_id"])
        per_hour_visits[bucket] += 1

    hourly = [[h, len(per_hour_people[h]), per_hour_visits[h]]
              for h in sorted(per_hour_people)]
    zones = [[z, per_zone_visits[z], round(per_zone_seconds[z] / per_zone_visits[z], 1),
              round(per_zone_seconds[z], 1)]
             for z in sorted(per_zone_visits)]

    busiest = max(hourly, key=lambda h: h[1]) if hourly else None
    return {"hourly": hourly, "zones": zones, "busiest": busiest,
            "undated": undated}


# ---------------------------------------------------------------------------
# writing


def ensure_out_dir(out_dir: Path) -> Path:
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportError(
            f"Could not create the output directory {out_dir}: {exc}\n"
            "  Pick a writable location with --out-dir."
        ) from exc
    if not os.access(out_dir, os.W_OK):
        raise ExportError(
            f"The output directory {out_dir} is not writable.\n"
            "  Pick a writable location with --out-dir."
        )
    return out_dir


def write_csv(path: Path, columns: Sequence[str], rows: Iterable[sqlite3.Row]) -> None:
    try:
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(columns)
            for r in rows:
                w.writerow([r[c] for c in columns])
    except OSError as exc:
        raise ExportError(f"Could not write {path}: {exc}") from exc


def write_xlsx(path: Path, columns: Sequence[str], rows: Sequence[sqlite3.Row],
               journeys: Sequence[list], summary: dict) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise ExportError(
            "The .xlsx export needs openpyxl.\n"
            "  pip install openpyxl\n"
            "  (the CSV above was still written, so nothing is lost)"
        ) from exc

    bold = Font(bold=True)
    wb = Workbook()

    def header(ws, values, row=1):
        ws.append(list(values))
        for cell in ws[row]:
            cell.font = bold

    # -- sheet 1: the raw rows ---------------------------------------------
    ws = wb.active
    ws.title = "Zone visits"
    header(ws, columns)
    for r in rows:
        ws.append([r[c] for c in columns])
    ws.freeze_panes = "A2"

    # -- sheet 2: one row per person ---------------------------------------
    ws = wb.create_sheet("Customer journeys")
    header(ws, ["person_id", "first_seen", "last_seen", "total_seconds",
                "zones_visited", "visits", "zone_path"])
    for j in journeys:
        ws.append(j)
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append([f"Note: {ESTIMATE_NOTE}"])

    # -- sheet 3: the rolled-up numbers ------------------------------------
    ws = wb.create_sheet("Footfall summary")
    ws.append(["Footfall summary"])
    ws["A1"].font = bold
    ws.append([f"Note: {ESTIMATE_NOTE}"])
    ws.append([])

    ws.append(["Busiest hour", (summary["busiest"][0] if summary["busiest"]
                                else "n/a")])
    ws["A4"].font = bold
    ws.append(["Estimated unique visitors that hour",
               (summary["busiest"][1] if summary["busiest"] else 0)])
    ws.append([])

    row_at = ws.max_row + 1
    header(ws, ["Hour", "Estimated unique visitors", "Visits"], row_at)
    for h in summary["hourly"]:
        ws.append(h)
    ws.append([])

    row_at = ws.max_row + 1
    header(ws, ["Zone", "Visits", "Average dwell (s)", "Total dwell (s)"], row_at)
    for z in summary["zones"]:
        ws.append(z)

    if summary["undated"]:
        ws.append([])
        ws.append([f"{summary['undated']} row(s) had an unreadable entered_at "
                   "and are excluded from the hourly table."])

    for sheet in wb.worksheets:
        _autosize(sheet)

    try:
        wb.save(path)
    except OSError as exc:
        raise ExportError(
            f"Could not write {path}: {exc}\n"
            "  If the file is open in Excel, close it and run again."
        ) from exc


def _autosize(ws, cap: int = 60) -> None:
    """Roughly fit column widths so the sheet is readable without dragging."""
    widths: dict = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column_letter] = min(
                cap, max(widths.get(cell.column_letter, 8), len(str(cell.value)) + 2))
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


# ---------------------------------------------------------------------------


def main(argv=None) -> int:
    p = argparse.ArgumentParser(
        description="Export the zone-visit database to a timestamped CSV + XLSX.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="The database is opened read-only, so this is safe to run while "
               "the tracker is writing to it.",
    )
    p.add_argument("--db", default=DEFAULT_DB,
                   help=f"path to the SQLite database (default: {DEFAULT_DB})")
    p.add_argument("--out-dir", default=DEFAULT_OUT_DIR,
                   help=f"directory for the export files (default: {DEFAULT_OUT_DIR})")
    a = p.parse_args(argv)

    db_path = Path(a.db).expanduser()
    out_dir = Path(a.out_dir).expanduser()

    try:
        conn = open_readonly(db_path)
        try:
            columns, rows = read_rows(conn)
        finally:
            conn.close()

        ensure_out_dir(out_dir)
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
        csv_path = out_dir / f"footfall_{stamp}.csv"
        xlsx_path = out_dir / f"footfall_{stamp}.xlsx"

        write_csv(csv_path, columns, rows)

        journeys = build_journeys(rows)
        summary = build_summary(rows)
        write_xlsx(xlsx_path, columns, rows, journeys, summary)
    except ExportError as exc:
        print(f"\n  {exc}\n", file=sys.stderr)
        return 1

    people = len({r["person_id"] for r in rows})
    busiest = summary["busiest"]
    print(f"\n  Exported {TABLE} from {db_path.resolve()}")
    print(f"    rows                       {len(rows)}")
    print(f"    distinct person_ids        {people}  (estimated unique visitors)")
    print(f"    zones                      {len(summary['zones'])}")
    if busiest:
        print(f"    busiest hour               {busiest[0]}  ({busiest[1]} est. visitors)")
    print("  Written:")
    print(f"    {csv_path.resolve()}")
    print(f"    {xlsx_path.resolve()}")
    print(f"\n  Note: {ESTIMATE_NOTE}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
